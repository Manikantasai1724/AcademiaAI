"""Multi-format academic document text extraction service.

Provides robust, format-specific extraction engines for PDF, DOCX, PPTX, and TXT
documents while preserving physical page and slide numbering for citations.
"""

from abc import ABC, abstractmethod
import hashlib
from pathlib import Path
from typing import Dict, List, Optional, Type

import pypdf
import docx
import pptx

from backend.app.documents.models import ExtractedDocument, ExtractedPage


class BaseExtractor(ABC):
    """Abstract base class for document format extractors."""

    @abstractmethod
    def extract(self, file_path: Path, filename: str, doc_id: str) -> ExtractedDocument:
        """Extract text and page-level metadata from a file on disk."""
        pass


class PDFExtractor(BaseExtractor):
    """Extracts text page-by-page from PDF academic documents using pypdf."""

    def extract(self, file_path: Path, filename: str, doc_id: str) -> ExtractedDocument:
        reader = pypdf.PdfReader(str(file_path))
        pages: List[ExtractedPage] = []

        total_pages = len(reader.pages)
        for idx, page in enumerate(reader.pages):
            page_number = idx + 1
            text = page.extract_text() or ""
            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    text=text,
                    metadata={"source_type": "page", "page_index": idx},
                )
            )

        if not pages:
            pages.append(
                ExtractedPage(
                    page_number=1,
                    text="",
                    metadata={"warning": "Empty PDF document"},
                )
            )

        return ExtractedDocument(
            document_id=doc_id,
            filename=filename,
            file_type="pdf",
            total_pages=len(pages),
            pages=pages,
        )


class DocxExtractor(BaseExtractor):
    """Extracts text, headings, and tables from DOCX files using python-docx."""

    def extract(self, file_path: Path, filename: str, doc_id: str) -> ExtractedDocument:
        doc = docx.Document(str(file_path))
        pages: List[ExtractedPage] = []
        current_page_number = 1
        current_page_text: List[str] = []
        current_section: str = "General"

        for para in doc.paragraphs:
            # Check for page break elements inside paragraph runs
            has_page_break = any(
                "lastRenderedPageBreak" in run._r.xml or "w:br" in run._r.xml and 'type="page"' in run._r.xml
                for run in para.runs
            )

            if has_page_break and current_page_text:
                pages.append(
                    ExtractedPage(
                        page_number=current_page_number,
                        text="\n".join(current_page_text).strip(),
                        metadata={"section": current_section},
                    )
                )
                current_page_number += 1
                current_page_text = []

            # Detect heading styles to annotate section context
            if para.style and para.style.name and "Heading" in para.style.name:
                current_section = para.text.strip() or current_section

            if para.text.strip():
                current_page_text.append(para.text.strip())

        # Extract table contents
        for table in doc.tables:
            table_rows: List[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                current_page_text.append("\n[Table Content]:\n" + "\n".join(table_rows))

        # Append final page
        pages.append(
            ExtractedPage(
                page_number=current_page_number,
                text="\n".join(current_page_text).strip(),
                metadata={"section": current_section},
            )
        )

        return ExtractedDocument(
            document_id=doc_id,
            filename=filename,
            file_type="docx",
            total_pages=len(pages),
            pages=pages,
        )


class PptxExtractor(BaseExtractor):
    """Extracts text, titles, and table contents slide-by-slide from PPTX presentations."""

    def extract(self, file_path: Path, filename: str, doc_id: str) -> ExtractedDocument:
        prs = pptx.Presentation(str(file_path))
        pages: List[ExtractedPage] = []

        for idx, slide in enumerate(prs.slides):
            slide_number = idx + 1
            slide_texts: List[str] = []
            slide_title: Optional[str] = None

            # Check for slide title
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()
                slide_texts.append(f"Slide Title: {slide_title}")

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue  # Already captured above

                # Extract text frames (bullet points, text boxes)
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)

                # Extract tables within slides
                if shape.has_table:
                    table_rows: List[str] = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            table_rows.append(" | ".join(row_cells))
                    if table_rows:
                        slide_texts.append("[Slide Table]:\n" + "\n".join(table_rows))

            # Capture speaker notes if present
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_texts.append(f"[Speaker Notes]: {notes_text}")

            pages.append(
                ExtractedPage(
                    page_number=slide_number,
                    text="\n".join(slide_texts).strip(),
                    metadata={
                        "source_type": "slide",
                        "slide_title": slide_title or f"Slide {slide_number}",
                    },
                )
            )

        if not pages:
            pages.append(
                ExtractedPage(
                    page_number=1,
                    text="",
                    metadata={"warning": "Empty presentation"},
                )
            )

        return ExtractedDocument(
            document_id=doc_id,
            filename=filename,
            file_type="pptx",
            total_pages=len(pages),
            pages=pages,
        )


class TextExtractor(BaseExtractor):
    """Extracts text from plain .txt files with multi-encoding fallback."""

    def extract(self, file_path: Path, filename: str, doc_id: str) -> ExtractedDocument:
        encodings = ["utf-8", "latin-1", "cp1252"]
        content: str = ""

        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    content = f.read()
                break
            except UnicodeDecodeError:
                continue

        pages = [
            ExtractedPage(
                page_number=1,
                text=content.strip(),
                metadata={"source_type": "plaintext"},
            )
        ]

        return ExtractedDocument(
            document_id=doc_id,
            filename=filename,
            file_type="txt",
            total_pages=1,
            pages=pages,
        )


class DocumentIngestionService:
    """Unified service dispatching document files to appropriate extractors."""

    _EXTRACTOR_REGISTRY: Dict[str, Type[BaseExtractor]] = {
        "pdf": PDFExtractor,
        "docx": DocxExtractor,
        "pptx": PptxExtractor,
        "txt": TextExtractor,
    }

    # Default file size limit: 50 MB
    MAX_FILE_SIZE_BYTES: int = 50 * 1024 * 1024

    @classmethod
    def generate_document_id(cls, file_path: Path, filename: str) -> str:
        """Compute deterministic SHA-256 hash from file bytes + filename."""
        hasher = hashlib.sha256()
        hasher.update(filename.encode("utf-8"))
        with open(file_path, "rb") as f:
            while chunk := f.read(8192):
                hasher.update(chunk)
        return f"doc_{hasher.hexdigest()[:16]}"

    @classmethod
    def extract_document(
        cls, file_path: Path, original_filename: Optional[str] = None
    ) -> ExtractedDocument:
        """Validate, detect format, and extract structured pages from a document."""
        if not file_path.exists():
            raise FileNotFoundError(f"Document file not found at: {file_path}")

        file_size = file_path.stat().st_size
        if file_size > cls.MAX_FILE_SIZE_BYTES:
            raise ValueError(
                f"File size ({file_size / (1024*1024):.2f} MB) exceeds maximum allowed limit "
                f"of {cls.MAX_FILE_SIZE_BYTES / (1024*1024):.0f} MB."
            )

        filename = original_filename or file_path.name
        file_ext = file_path.suffix.lower().lstrip(".")

        if file_ext not in cls._EXTRACTOR_REGISTRY:
            supported = ", ".join(cls._EXTRACTOR_REGISTRY.keys())
            raise ValueError(
                f"Unsupported file format '.{file_ext}'. Supported formats are: {supported}."
            )

        doc_id = cls.generate_document_id(file_path, filename)
        extractor_cls = cls._EXTRACTOR_REGISTRY[file_ext]
        extractor = extractor_cls()

        return extractor.extract(file_path=file_path, filename=filename, doc_id=doc_id)
